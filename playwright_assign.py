import openpyxl
import json
import random
from datetime import datetime
from pathlib import Path
from playwright.sync_api import sync_playwright


# ============================================================
# CONFIGURATION
# ============================================================

INPUT_FILE = "contacts.xlsx"

REPORT_DATE = datetime.now().strftime("%Y-%m-%d")

JSON_FILE = f"whatsapp_report_{REPORT_DATE}.json"
EXCEL_FILE = f"whatsapp_report_{REPORT_DATE}.xlsx"

SCREENSHOT_DIR = Path("screenshots")
SCREENSHOT_DIR.mkdir(exist_ok=True)

PROFILE_DIR = "whatsapp_profile"


# ============================================================
# LOAD CONTACTS FROM EXCEL
# ============================================================

workbook = openpyxl.load_workbook(INPUT_FILE)
worksheet = workbook.active

contacts = []

for row in worksheet.iter_rows(min_row=2, values_only=True):

    name = row[0]
    phone = row[1]
    message = row[2] if len(row) > 2 else None

    if not name or not phone:
        continue

    contacts.append({
        "name": str(name).strip(),
        "phone": str(phone).strip(),
        "message": str(message).strip()
        if message else "Hello {name}, this is a test message from my Playwright automation bot."
    })


print(f"Contacts loaded: {len(contacts)}")


# ============================================================
# RESULT STORAGE
# ============================================================

results = []


# ============================================================
# PLAYWRIGHT
# ============================================================

with sync_playwright() as p:

    context = p.chromium.launch_persistent_context(
        PROFILE_DIR,
        headless=False
    )

    page = context.pages[0] if context.pages else context.new_page()

    page.goto("https://web.whatsapp.com")

    input(
        "Press ENTER when WhatsApp Web is fully loaded "
        "and logged in..."
    )


    # ========================================================
    # PROCESS EACH CONTACT
    # ========================================================

    for contact in contacts:

        name = contact["name"]
        phone = contact["phone"]

        personalized_message = contact["message"].replace(
            "{name}",
            name
        )

        print("\n" + "=" * 60)
        print(f"Processing: {name}")
        print(f"Phone: {phone}")
        print(f"Message: {personalized_message}")
        print("=" * 60)

        result = {
            "name": name,
            "phone": phone,
            "message": personalized_message,
            "status": "Failed",
            "screenshot": "",
            "last_3_messages": [],
            "error": ""
        }

        try:

            # ------------------------------------------------
            # SEARCH CONTACT
            # ------------------------------------------------

            search_box = page.get_by_role(
                "textbox",
                name="Search or start a new chat"
            )

            search_box.wait_for(
                state="visible",
                timeout=30000
            )

            search_box.fill(phone)

            page.wait_for_timeout(3000)

            print("Contact search completed.")


            # ------------------------------------------------
            # OPEN SEARCH RESULT
            # ------------------------------------------------

            # Try phone number first
            opened = False

            try:

                phone_result = page.get_by_text(
                    phone,
                    exact=False
                ).first

                phone_result.wait_for(
                    state="visible",
                    timeout=5000
                )

                phone_result.click()

                opened = True

                print("Search result opened using phone.")

            except Exception:
                pass


            # Try contact name
            if not opened:

                try:

                    name_result = page.get_by_text(
                        name,
                        exact=True
                    ).first

                    name_result.wait_for(
                        state="visible",
                        timeout=5000
                    )

                    name_result.click()

                    opened = True

                    print("Search result opened using name.")

                except Exception:
                    pass


            # Keyboard fallback
            if not opened:

                print("Trying keyboard selection...")

                page.keyboard.press("ArrowDown")

                page.wait_for_timeout(1000)

                page.keyboard.press("Enter")

                page.wait_for_timeout(3000)


            # ------------------------------------------------
            # FIND MESSAGE COMPOSER
            # ------------------------------------------------

            print("Looking for message box...")

            message_box = page.locator(
                '[aria-label^="Type a message to "]'
            ).first

            message_box.wait_for(
                state="visible",
                timeout=15000
            )

            print("Message box found.")


            # ------------------------------------------------
            # RANDOM DELAY
            # ------------------------------------------------

            delay = random.uniform(2, 5)

            print(f"Waiting {delay:.1f} seconds...")

            page.wait_for_timeout(
                int(delay * 1000)
            )


            # ------------------------------------------------
            # TYPE MESSAGE
            # ------------------------------------------------

            message_box.click()

            page.keyboard.type(
                personalized_message
            )

            page.wait_for_timeout(1000)

            print("Message entered.")


            # ------------------------------------------------
            # SEND MESSAGE
            # ------------------------------------------------

            page.keyboard.press("Enter")

            page.wait_for_timeout(3000)

            print("Message sent.")

            result["status"] = "Sent"


            # ------------------------------------------------
            # SCREENSHOT
            # ------------------------------------------------

            safe_name = "".join(
                c if c.isalnum() else "_"
                for c in name
            )

            screenshot_file = (
                SCREENSHOT_DIR
                / f"{safe_name}_{REPORT_DATE}.png"
            )

            page.screenshot(
                path=str(screenshot_file),
                full_page=True
            )

            result["screenshot"] = str(
                screenshot_file
            )

            print(
                f"Screenshot saved: {screenshot_file}"
            )


            # ------------------------------------------------
            # EXTRACT LAST 3 MESSAGES
            # ------------------------------------------------

            print("Extracting last 3 messages...")

            message_elements = page.locator(
                '[data-pre-plain-text]'
            )

            message_count = message_elements.count()

            extracted_messages = []

            start_index = max(
                0,
                message_count - 3
            )

            for i in range(
                start_index,
                message_count
            ):

                try:

                    text = (
                        message_elements
                        .nth(i)
                        .inner_text()
                        .strip()
                    )

                    if text:
                        extracted_messages.append(text)

                except Exception:
                    continue

            result["last_3_messages"] = (
                extracted_messages[-3:]
            )

            print(
                f"Messages extracted: "
                f"{len(result['last_3_messages'])}"
            )


        except Exception as e:

            result["status"] = "Failed"

            result["error"] = str(e)

            print(
                f"ERROR processing {name}: {e}"
            )


        results.append(result)


        # ----------------------------------------------------
        # DELAY BEFORE NEXT CONTACT
        # ----------------------------------------------------

        if contact != contacts[-1]:

            delay = random.uniform(2, 5)

            print(
                f"Waiting {delay:.1f} seconds "
                "before next contact..."
            )

            page.wait_for_timeout(
                int(delay * 1000)
            )


    # ========================================================
    # SAVE JSON REPORT
    # ========================================================

    with open(
        JSON_FILE,
        "w",
        encoding="utf-8"
    ) as json_file:

        json.dump(
            results,
            json_file,
            indent=4,
            ensure_ascii=False
        )

    print(f"\nJSON report saved: {JSON_FILE}")


    # ========================================================
    # SAVE EXCEL REPORT
    # ========================================================

    report_workbook = openpyxl.Workbook()

    report_sheet = report_workbook.active

    report_sheet.title = "WhatsApp Report"

    headers = [
        "Name",
        "Phone",
        "Message",
        "Status",
        "Screenshot",
        "Last 3 Messages",
        "Error"
    ]

    report_sheet.append(headers)

    for item in results:

        report_sheet.append([
            item["name"],
            item["phone"],
            item["message"],
            item["status"],
            item["screenshot"],
            " | ".join(
                item["last_3_messages"]
            ),
            item["error"]
        ])


    report_workbook.save(EXCEL_FILE)

    print(
        f"Excel report saved: {EXCEL_FILE}"
    )


    # ========================================================
    # FINISH
    # ========================================================

    print("\n" + "=" * 60)
    print("ASSIGNMENT 2 COMPLETED")
    print("=" * 60)

    print(f"JSON : {JSON_FILE}")
    print(f"Excel: {EXCEL_FILE}")
    print(f"Screenshots: {SCREENSHOT_DIR}")

    input("\nPress ENTER to close WhatsApp...")

    context.close()